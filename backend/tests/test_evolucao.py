"""Cartões editáveis, categorias (CRUD/mesclagem), orçamento, lembretes, exportação e competência de parcelados."""
import io
from datetime import date, timedelta

import pytest
from openpyxl import load_workbook

from app import gemini
from conftest import categoria_por_codigo, conta
from test_captura import foto, resposta
from test_cartao import cenario, compra  # noqa: F401  (fixture)


def novo_cartao(p, cc=None, fech=10, venc=20, final="1234"):
    r = p.post("/api/cartoes", json={"nome": "Visa", "dia_fechamento": fech, "dia_vencimento": venc,
                                     "conta_pagamento_id": cc["id"] if cc else None, "final_principal": final})
    assert r.status_code == 201, r.text
    return r.json()


def lanca_conta(p, cc, cat, valor, data, tipo="despesa", **extra):
    r = p.post("/api/transacoes", json={"tipo": tipo, "valor_centavos": valor, "data_competencia": data, "conta_id": cc["id"],
                                        "categoria_id": cat["id"], **extra})
    assert r.status_code == 201, r.text
    return r.json()[0]


# ---------- cartões ----------
def test_editar_e_excluir_conta_de_cartao(nova_pessoa):
    a, outro = nova_pessoa("dono"), nova_pessoa("outro")
    cc = conta(a)
    k = novo_cartao(a, cc)
    r = a.patch(f"/api/cartoes/{k['id']}", json={"nome": "Visa Gold", "limite_centavos": 500000, "conta_pagamento_id": None})
    assert r.status_code == 200 and r.json()["nome"] == "Visa Gold" and r.json()["conta_pagamento_id"] is None
    assert outro.patch(f"/api/cartoes/{k['id']}", json={"nome": "x"}).status_code == 404
    assert outro.delete(f"/api/cartoes/{k['id']}").status_code == 404
    cat = categoria_por_codigo(a, "2020.01")
    pid = k["plasticos"][0]["id"]
    assert compra(a, pid, cat["id"], 1000, "2026-09-05").status_code == 201
    d = a.delete(f"/api/cartoes/{k['id']}")
    assert d.status_code == 409 and "compra" in d.json()["detail"]              # preserva histórico por padrão
    assert a.patch(f"/api/cartoes/{k['id']}", json={"inativo": True}).json()["inativo"] is True
    d = a.delete(f"/api/cartoes/{k['id']}?com_historico=true")
    assert d.status_code == 200 and d.json()["compras_excluidas"] == 1
    assert a.get("/api/cartoes").json() == [] and a.get("/api/transacoes").json() == []
    k2 = novo_cartao(a, cc, final="9999")                                       # sem compras: exclui direto
    assert a.delete(f"/api/cartoes/{k2['id']}").status_code == 200


def test_tipos_e_cartao_principal_unico(nova_pessoa):
    a = nova_pessoa("dono")
    k = novo_cartao(a, conta(a))
    princ = k["plasticos"][0]
    assert princ["principal"] and princ["tipo"] == "plastico"
    virt = a.post(f"/api/cartoes/{k['id']}/plasticos", json={"final": "1111", "rotulo": "Compras online", "tipo": "virtual"}).json()
    ad = a.post(f"/api/cartoes/{k['id']}/plasticos", json={"final": "2222", "rotulo": "Maria"}).json()
    assert virt["tipo"] == "virtual" and ad["tipo"] == "plastico" and not virt["principal"]
    assert a.patch(f"/api/plasticos/{virt['id']}", json={"principal": True}).status_code == 422          # virtual nunca é principal
    assert a.patch(f"/api/plasticos/{princ['id']}", json={"tipo": "virtual"}).status_code == 422
    assert a.patch(f"/api/plasticos/{princ['id']}", json={"ativo": False}).status_code == 422
    assert a.delete(f"/api/plasticos/{princ['id']}").status_code == 422
    r = a.patch(f"/api/plasticos/{ad['id']}", json={"principal": True, "rotulo": "Titular"})
    assert r.status_code == 200 and r.json()["principal"] is True
    cartao = a.get("/api/cartoes").json()[0]
    assert [p["final"] for p in cartao["plasticos"] if p["principal"]] == ["2222"]                        # só um principal
    assert a.patch(f"/api/plasticos/{virt['id']}", json={"final": "3333", "rotulo": "Virtual 2"}).json()["final"] == "3333"
    cat = categoria_por_codigo(a, "2020.01")
    assert compra(a, virt["id"], cat["id"], 500, "2026-09-05").status_code == 201
    assert a.delete(f"/api/plasticos/{virt['id']}").status_code == 409                                    # tem compras
    assert a.delete(f"/api/plasticos/{princ['id']}").status_code == 200                                   # não é mais o principal


def test_mudar_fechamento_recalcula_faturas_abertas(nova_pessoa):
    a = nova_pessoa("dono")
    k = novo_cartao(a, conta(a), fech=20, venc=25)
    cat = categoria_por_codigo(a, "2020.01")
    hoje = date.today().isoformat()
    t = compra(a, k["plasticos"][0]["id"], cat["id"], 1000, hoje).json()[0]
    assert t["data_caixa"].endswith("-25")
    r = a.patch(f"/api/cartoes/{k['id']}", json={"dia_fechamento": 21, "dia_vencimento": 26})
    assert r.status_code == 200 and r.json()["faturas_recalculadas"] == 1
    f = a.get(f"/api/cartoes/{k['id']}/faturas").json()[0]
    assert f["data_fechamento"].endswith("-21") and f["data_vencimento"].endswith("-26")
    assert a.get("/api/transacoes").json()[0]["data_caixa"].endswith("-26")     # a data de caixa do item acompanha
    assert a.patch(f"/api/cartoes/{k['id']}", json={"nome": "Só nome"}).json()["faturas_recalculadas"] == 0


def test_parcelado_com_competencia_da_compra_inteira(cenario):
    a, s = cenario["a"], cenario
    cat = s["cat_a"]["id"]
    r = compra(a, s["principal"]["id"], cat, 30000, "2026-09-05", parcelas=3, competencia_parcelas="compra")
    assert r.status_code == 201, r.text
    ts = r.json()
    assert [t["data_competencia"] for t in ts] == ["2026-09-05"] * 3
    assert [t["data_caixa"] for t in ts] == ["2026-09-20", "2026-10-20", "2026-11-20"]      # o caixa segue as faturas
    assert {t["estado"] for t in ts} == {"confirmado"}
    assert a.get("/api/resumo/mensal?mes=2026-09").json()["despesas"] == -30000
    # padrão: cada parcela no próprio mês
    r = compra(a, s["principal"]["id"], cat, 9000, "2026-09-05", parcelas=3).json()
    assert [t["data_competencia"] for t in r] == ["2026-09-05", "2026-10-05", "2026-11-05"]
    assert [t["estado"] for t in r] == ["confirmado", "previsto", "previsto"]


# ---------- categorias ----------
def test_categorias_crud_mover_e_excluir(nova_pessoa):
    a, b = nova_pessoa("dono"), nova_pessoa("outro")
    g = a.post("/api/categorias", json={"nome": "Pets", "tipo": "despesa"})
    assert g.status_code == 201 and g.json()["pai_id"] is None
    assert a.post("/api/categorias", json={"nome": "Sem tipo"}).status_code == 422
    assert a.post("/api/categorias", json={"nome": "pets", "tipo": "despesa"}).status_code == 409        # duplicado (sem diferenciar maiúsculas)
    g = g.json()
    sub = a.post("/api/categorias", json={"nome": "Ração", "pai_id": g["id"]}).json()
    assert sub["tipo"] == "despesa" and sub["pai_id"] == g["id"]
    assert a.patch(f"/api/categorias/{sub['id']}", json={"nome": "Ração e petiscos"}).json()["nome"] == "Ração e petiscos"
    saude = categoria_por_codigo(a, "2070.07")                                                          # Exames
    grupo_receita = categoria_por_codigo(a, "1000")
    assert a.patch(f"/api/categorias/{sub['id']}", json={"pai_id": grupo_receita["id"]}).status_code == 422   # tipos diferentes
    assert a.patch(f"/api/categorias/{sub['id']}", json={"pai_id": saude["pai_id"]}).json()["pai_id"] == saude["pai_id"]
    assert a.patch(f"/api/categorias/{g['id']}", json={"pai_id": saude["pai_id"]}).status_code == 422    # grupo não vira subcategoria
    # de outro usuário: invisível/intocável
    assert b.patch(f"/api/categorias/{sub['id']}", json={"nome": "x"}).status_code == 404
    assert b.delete(f"/api/categorias/{sub['id']}").status_code == 404
    # grupo com filhos não exclui; em uso não exclui
    assert a.delete(f"/api/categorias/{saude['pai_id']}").status_code == 409
    cc = conta(a)
    lanca_conta(a, cc, sub, 1000, "2026-09-05")
    assert a.delete(f"/api/categorias/{sub['id']}").status_code == 409
    assert a.patch(f"/api/categorias/{sub['id']}", json={"ativa": False}).json()["ativa"] is False
    assert a.delete(f"/api/categorias/{g['id']}").status_code == 200                                    # grupo agora vazio


def test_mesclar_categorias_move_tudo_e_preserva_codigo_da_ia(cenario, monkeypatch):
    a, b, s = cenario["a"], cenario["b"], cenario
    cc = s["conta"]
    origem = categoria_por_codigo(a, "2020.02")     # Padaria e dia a dia
    destino = categoria_por_codigo(a, "2020.01")    # Supermercado
    t = lanca_conta(a, cc, origem, 2500, "2026-09-05")
    tp = compra(b, s["adicional"]["id"], origem["id"], 700, "2026-09-06").json()[0]       # compra do portador na categoria do dono
    a.put("/api/orcamento", json={"itens": [{"categoria_id": origem["id"], "valor_centavos": 10000},
                                            {"categoria_id": destino["id"], "valor_centavos": 50000}]})
    assert a.post(f"/api/categorias/{origem['id']}/mesclar", json={"destino_id": origem["id"]}).status_code == 422
    assert a.post(f"/api/categorias/{origem['id']}/mesclar", json={"destino_id": categoria_por_codigo(a, "1000.01")["id"]}).status_code == 422
    r = a.post(f"/api/categorias/{origem['id']}/mesclar", json={"destino_id": destino["id"]})
    assert r.status_code == 200 and r.json()["lancamentos_movidos"] == 2
    ids = {x["id"]: x for x in a.get("/api/transacoes").json()}
    assert ids[t["id"]]["categoria_id"] == destino["id"] and ids[tp["id"]]["categoria_id"] == destino["id"]
    assert origem["id"] not in {c["id"] for c in a.get("/api/categorias").json()}
    orc = a.get("/api/orcamento?mes=2026-09").json()
    linha = next(c for g in orc["despesas"]["grupos"] for c in g["categorias"] if c["id"] == destino["id"])
    assert linha["orcado"] == 60000 and linha["realizado"] == 3200                                  # orçamentos e gastos somados
    # a IA continua mapeando o código antigo (2020.02) para a categoria de destino
    monkeypatch.setattr(gemini, "MOCK_RESPOSTA", resposta(categoria_codigo="2020.02", estabelecimento="OUTRO LUGAR", final_cartao=""))
    sug = a.post("/api/capturas", files=foto()).json()["sugestao"]
    assert sug["categoria_id"] == destino["id"]


# ---------- orçamento ----------
def test_orcamento_previsto_x_real_por_competencia(nova_pessoa):
    a = nova_pessoa("dono")
    cc = conta(a)
    mercado, luz = categoria_por_codigo(a, "2020.01"), categoria_por_codigo(a, "2010.05")
    salario = categoria_por_codigo(a, "1000.01")
    lanca_conta(a, cc, mercado, 30000, "2026-09-03")
    lanca_conta(a, cc, mercado, 12000, "2026-09-20", estado="previsto")
    lanca_conta(a, cc, mercado, 99900, "2026-08-31")                                  # outra competência
    lanca_conta(a, cc, luz, 25000, "2026-09-10")
    lanca_conta(a, cc, salario, 800000, "2026-09-05", tipo="receita")
    a.put("/api/orcamento", json={"itens": [{"categoria_id": mercado["id"], "valor_centavos": 50000},
                                            {"categoria_id": luz["id"], "valor_centavos": 20000},
                                            {"categoria_id": salario["id"], "valor_centavos": 900000},
                                            {"categoria_id": mercado["id"], "valor_centavos": 70000, "mes": "2026-09"}]})
    o = a.get("/api/orcamento?mes=2026-09").json()
    assert o["pode_editar"] is True
    linha = {c["id"]: c for tipo in ("despesas", "receitas") for g in o[tipo]["grupos"] for c in g["categorias"]}
    m = linha[mercado["id"]]
    assert (m["orcado"], m["realizado"], m["previsto"], m["especifico"]) == (70000, 30000, 12000, True)     # ajuste do mês vale
    assert linha[luz["id"]]["realizado"] == 25000 and linha[luz["id"]]["orcado"] == 20000
    assert o["receitas"]["realizado"] == 800000 and o["receitas"]["orcado"] == 900000
    assert o["despesas"]["orcado"] == 90000 and o["despesas"]["realizado"] == 55000 and o["despesas"]["previsto"] == 12000
    ago = a.get("/api/orcamento?mes=2026-08").json()
    assert next(c for g in ago["despesas"]["grupos"] for c in g["categorias"] if c["id"] == mercado["id"])["orcado"] == 50000   # volta ao padrão
    assert a.delete(f"/api/orcamento?categoria_id={mercado['id']}&mes=2026-09").status_code == 200
    assert a.delete(f"/api/orcamento?categoria_id={mercado['id']}&mes=2026-09").status_code == 404
    assert a.put("/api/orcamento", json={"itens": [{"categoria_id": mercado["id"], "valor_centavos": -1}]}).status_code == 422


def test_orcamento_media_dos_ultimos_meses(nova_pessoa):
    a = nova_pessoa("dono")
    cc = conta(a)
    cat = categoria_por_codigo(a, "2020.01")
    hoje = date.today().replace(day=1)
    for i in (1, 2, 3):
        lanca_conta(a, cc, cat, 30000 * i, (hoje - timedelta(days=31 * i - 1)).replace(day=15).isoformat())
    m = a.get("/api/orcamento/media?meses=3").json()
    assert m == [{"categoria_id": cat["id"], "valor_centavos": 60000}]              # (300+600+900)/3


def test_orcamento_consulta_de_quem_compartilha_sem_vazar(nova_pessoa):
    a, b, estranho = nova_pessoa("dono"), nova_pessoa("parceiro"), nova_pessoa("estranho")
    conjunta, privada = conta(a, "Conjunta"), conta(a, "Privada")
    cat = categoria_por_codigo(a, "2020.01")
    lanca_conta(a, conjunta, cat, 10000, "2026-09-05")
    lanca_conta(a, privada, cat, 77700, "2026-09-06")
    a.put("/api/orcamento", json={"itens": [{"categoria_id": cat["id"], "valor_centavos": 100000}]})
    v = a.post(f"/api/contas/{conjunta['id']}/convites", json={"email": b.email, "papel": "leitor"}).json()
    assert b.post(f"/api/convites/{v['id']}/aceitar").status_code == 200
    o = b.get(f"/api/orcamento?mes=2026-09&dono_id={a.id}")
    assert o.status_code == 200 and o.json()["pode_editar"] is False and o.json()["dono_nome"]
    linha = next(c for g in o.json()["despesas"]["grupos"] for c in g["categorias"] if c["id"] == cat["id"])
    assert linha["orcado"] == 100000 and linha["realizado"] == 10000               # só o gasto da conta compartilhada
    assert {d["id"] for d in b.get("/api/orcamento/donos").json()} == {a.id, b.id}
    assert b.put("/api/orcamento", json={"itens": [{"categoria_id": cat["id"], "valor_centavos": 1}]}).status_code == 422
    assert b.delete(f"/api/orcamento?categoria_id={cat['id']}").status_code == 404
    assert estranho.get(f"/api/orcamento?mes=2026-09&dono_id={a.id}").status_code == 404
    assert estranho.get("/api/orcamento/donos").json() == [{"id": estranho.id, "nome": estranho.get("/api/eu").json()["nome"]}]


# ---------- lembretes ----------
def test_lembretes_previstos_e_faturas(cenario):
    a, b, s = cenario["a"], cenario["b"], cenario
    cat = s["cat_a"]
    hoje = date.today()
    lanca_conta(a, s["conta"], cat, 5000, (hoje - timedelta(days=2)).isoformat(), estado="previsto", descricao="Atrasada")
    lanca_conta(a, s["conta"], cat, 7000, (hoje + timedelta(days=3)).isoformat(), estado="previsto", descricao="Em 3 dias")
    lanca_conta(a, s["conta"], cat, 9000, (hoje + timedelta(days=20)).isoformat(), estado="previsto", descricao="Em 20 dias")
    lanca_conta(a, s["conta"], cat, 1234, hoje.isoformat(), descricao="Já pago")                        # confirmado: não é lembrete
    compra(a, s["principal"]["id"], cat["id"], 4000, hoje.isoformat())
    r = a.get("/api/lembretes?dias=7").json()
    descr = [i.get("descricao") or i["tipo"] for i in r["itens"]]
    assert "Atrasada" in descr and "Em 3 dias" in descr and "Em 20 dias" not in descr and "Já pago" not in descr
    atrasado = next(i for i in r["itens"] if i.get("descricao") == "Atrasada")
    assert atrasado["atrasado"] is True and r["atrasados"] >= 1
    assert r["saidas"] <= -12000
    assert len(a.get("/api/lembretes?dias=30").json()["itens"]) > len(r["itens"])
    assert any(i["tipo"] == "fatura" for i in a.get("/api/lembretes?dias=60").json()["itens"])
    # o portador não enxerga contas nem faturas do dono
    assert b.get("/api/lembretes?dias=60").json()["itens"] == []


# ---------- exportação ----------
def _abrir(resp):
    assert resp.status_code == 200, resp.text
    assert "attachment" in resp.headers["content-disposition"] and ".xlsx" in resp.headers["content-disposition"]
    return load_workbook(io.BytesIO(resp.content))


def test_exportar_excel_lancamentos_e_faturas(cenario):
    a, b, s = cenario["a"], cenario["b"], cenario
    cat = s["cat_a"]
    lanca_conta(a, s["conta"], cat, 12345, "2026-09-05", descricao="=1+1", favorecido_nome="Mercado")
    compra(a, s["principal"]["id"], cat["id"], 3000, "2026-09-06", favorecido_nome="Loja")
    compra(b, s["adicional"]["id"], cat["id"], 4590, "2026-09-07", favorecido_nome="Padaria")
    wb = _abrir(a.get("/api/exportar?de=2026-09-01&ate=2026-09-30"))
    assert wb.sheetnames == ["Lançamentos", "Faturas", "Itens das faturas"]
    linhas = list(wb["Lançamentos"].iter_rows(min_row=2, values_only=True))
    assert len(linhas) == 3
    conta_row = next(l for l in linhas if l[6] == "Corrente")
    assert conta_row[5] == -123.45 and conta_row[2] is None
    celula = next(c for row in wb["Lançamentos"].iter_rows(min_row=2) for c in row if c.value == "=1+1")
    assert celula.data_type == "s"                                            # texto, não fórmula
    fat = list(wb["Faturas"].iter_rows(min_row=2, values_only=True))
    assert len(fat) == 1 and fat[0][0] == "Visa" and fat[0][5] == 75.90 and fat[0][6] == 2
    itens = list(wb["Itens das faturas"].iter_rows(min_row=2, values_only=True))
    assert sorted(i[11] for i in itens) == [30.00, 45.90] and {i[5] for i in itens} == {None, b.get("/api/eu").json()["nome"]}
    # portador: só as compras dele e nenhuma fatura
    lp = list(_abrir(b.get("/api/exportar?de=2026-09-01&ate=2026-10-31&base=caixa"))["Lançamentos"].iter_rows(min_row=2, values_only=True))
    assert len(lp) == 1 and lp[0][5] == -45.90
    assert list(_abrir(b.get("/api/exportar?de=2026-09-01&ate=2026-10-31"))["Faturas"].iter_rows(min_row=2, values_only=True)) == []
    assert a.get("/api/exportar?de=2026-10-01&ate=2026-09-01").status_code == 422


# ---------- orçamento com vigência ----------
def _orc(a, cat, mes):
    o = a.get(f"/api/orcamento?mes={mes}").json()
    for bloco in ("despesas", "receitas"):
        for g in o[bloco]["grupos"]:
            for c in g["categorias"]:
                if c["id"] == cat["id"]:
                    return c
    raise AssertionError("categoria não listada")


def test_orcamento_vigencia_continuo_ocorrencias_e_meses(nova_pessoa):
    a = nova_pessoa("vig")
    mercado, luz, ipva = categoria_por_codigo(a, "2020.01"), categoria_por_codigo(a, "2010.05"), categoria_por_codigo(a, "2090.03")
    put = lambda **k: a.put("/api/orcamento", json={"itens": [k]})
    # 1) sem término a partir de um mês (com término opcional)
    assert put(categoria_id=mercado["id"], valor_centavos=50000, modo="continuo", inicio="2026-03").status_code == 200
    assert _orc(a, mercado, "2026-02")["orcado"] == 0
    assert _orc(a, mercado, "2026-03")["orcado"] == 50000 and _orc(a, mercado, "2031-01")["orcado"] == 50000
    # uma regra mais recente substitui a anterior daquele mês em diante
    assert put(categoria_id=mercado["id"], valor_centavos=60000, modo="continuo", inicio="2026-07", fim="2026-09").status_code == 200
    assert _orc(a, mercado, "2026-06")["orcado"] == 50000 and _orc(a, mercado, "2026-08")["orcado"] == 60000
    assert _orc(a, mercado, "2026-10")["orcado"] == 50000                         # passou do término: volta a valer a regra base
    # 2) número de ocorrências
    assert put(categoria_id=luz["id"], valor_centavos=20000, modo="ocorrencias", inicio="2026-09", ocorrencias=3).status_code == 200
    assert [_orc(a, luz, m)["orcado"] for m in ("2026-08", "2026-09", "2026-11", "2026-12")] == [0, 20000, 20000, 0]
    # 3) meses específicos do ano (repete todo ano)
    assert put(categoria_id=ipva["id"], valor_centavos=300000, modo="meses", inicio="2026-01", meses=[1, 7]).status_code == 200
    assert [_orc(a, ipva, m)["orcado"] for m in ("2026-01", "2026-02", "2026-07", "2027-07", "2027-08")] == [300000, 0, 300000, 300000, 0]
    # ajuste de um mês vence sempre
    assert put(categoria_id=ipva["id"], valor_centavos=1, mes="2027-07").status_code == 200
    assert _orc(a, ipva, "2027-07")["orcado"] == 1 and _orc(a, ipva, "2027-07")["especifico"] is True
    # regras listáveis e removíveis uma a uma
    regras = a.get(f"/api/orcamento/regras?categoria_id={mercado['id']}").json()
    assert {r["modo"] for r in regras} == {"continuo"} and len(regras) == 2
    alvo = next(r for r in regras if r["inicio"] == "2026-07")
    assert a.delete(f"/api/orcamento?regra_id={alvo['id']}").status_code == 200
    assert _orc(a, mercado, "2026-08")["orcado"] == 50000
    # validações
    assert put(categoria_id=luz["id"], valor_centavos=1, modo="ocorrencias", inicio="2026-09").status_code == 422
    assert put(categoria_id=luz["id"], valor_centavos=1, modo="meses", inicio="2026-09", meses=[13]).status_code == 422
    assert put(categoria_id=luz["id"], valor_centavos=1, modo="ocorrencias", inicio="2026-09", ocorrencias=2, fim="2026-12").status_code == 422
    assert put(categoria_id=luz["id"], valor_centavos=1, modo="continuo", inicio="2026-09", fim="2026-01").status_code == 422
    assert put(categoria_id=luz["id"], valor_centavos=1, modo="meses", meses=[1]).status_code == 422       # precisa de início


def test_lembretes_mostra_fatura_alem_do_periodo(cenario):
    a, s = cenario["a"], cenario
    compra(a, s["principal"]["id"], s["cat_a"]["id"], 4000, date.today().isoformat())
    curto = a.get("/api/lembretes?dias=0").json()
    fat = [i for i in curto["itens"] if i["tipo"] == "fatura"]
    assert fat and all(i["valor_centavos"] for i in fat)
    alem = [i for i in fat if i.get("alem_periodo")]
    if alem:                                      # fatura fora do período aparece, mas não entra nos totais
        assert curto["saidas"] == sum(i["valor_centavos"] for i in curto["itens"] if i["valor_centavos"] < 0 and not i.get("alem_periodo"))
