import io
from datetime import date

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..deps import get_db

router = APIRouter(prefix="/api")
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_LINHAS = 100_000




def _escrever(ws, colunas, linhas):
    """colunas: [(titulo, largura, formato)]; linhas: listas de valores. Datas viram datas; strings perigosas viram texto."""
    ws.append([c[0] for c in colunas])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0F766E")
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for i, (_, larg, fmt) in enumerate(colunas, start=1):
        ws.column_dimensions[get_column_letter(i)].width = larg
    for lin in linhas:
        ws.append(lin)
        for j, cel in enumerate(ws[ws.max_row]):
            fmt = colunas[j][2]
            if fmt:
                cel.number_format = fmt
            if isinstance(cel.value, str) and cel.value[:1] in ("=", "+", "-", "@"):
                cel.data_type = "s"
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions


def _reais(c):
    return None if c is None else round(c / 100, 2)


def _cartao(nome, final, tipo):
    return f"{nome} ·· {final}{' (virtual)' if tipo == 'virtual' else ''}" if nome else None


@router.get("/exportar")
def exportar(de: date, ate: date, base: str = Query("competencia", pattern="^(competencia|caixa)$"), cur=Depends(get_db)):
    """Planilha .xlsx com lançamentos, faturas de cartão e itens das faturas do período (só o que a pessoa pode ver)."""
    if ate < de:
        raise HTTPException(422, "a data final é anterior à inicial")
    col = "t.data_competencia" if base == "competencia" else "t.data_caixa"
    lancs = cur.execute(
        f"""SELECT t.data_competencia, t.data_caixa, t.data_compra, t.tipo, t.estado, t.valor_centavos, c.nome AS conta,
                   k.nome AS cartao, pl.final, pl.tipo AS tipo_cartao, pu.nome AS portador, g.nome AS grupo, cat.nome AS categoria,
                   f.nome AS favorecido, t.descricao, t.forma_pagamento, t.numero_parcela, t.total_parcelas, u.nome AS lancado_por, t.origem
            FROM transacao t LEFT JOIN conta c ON c.id = t.conta_id LEFT JOIN plastico pl ON pl.id = t.plastico_id
            LEFT JOIN cartao k ON k.id = pl.cartao_id LEFT JOIN usuario pu ON pu.id = pl.portador_id
            LEFT JOIN categoria cat ON cat.id = t.categoria_id LEFT JOIN categoria g ON g.id = cat.pai_id
            LEFT JOIN favorecido f ON f.id = t.favorecido_id LEFT JOIN usuario u ON u.id = t.criado_por
            WHERE {col} BETWEEN %s AND %s ORDER BY {col}, t.criado_em LIMIT %s""", (de, ate, MAX_LINHAS + 1)).fetchall()
    if len(lancs) > MAX_LINHAS:
        raise HTTPException(422, "período muito grande: reduza o intervalo")
    faturas = cur.execute(
        """SELECT k.nome AS cartao, f.mes_referencia, f.data_fechamento, f.data_vencimento, f.status,
                  COALESCE(SUM(t.valor_centavos), 0)::bigint AS total, COUNT(t.id)::int AS itens
           FROM fatura f JOIN cartao k ON k.id = f.cartao_id LEFT JOIN transacao t ON t.fatura_id = f.id
           WHERE f.data_vencimento BETWEEN %s AND %s GROUP BY f.id, k.nome ORDER BY f.data_vencimento, k.nome""", (de, ate)).fetchall()
    itens = cur.execute(
        """SELECT k.nome AS cartao, f.data_vencimento, f.status, t.data_compra, t.data_competencia, pl.final, pl.tipo AS tipo_cartao,
                  pu.nome AS portador, g.nome AS grupo, cat.nome AS categoria, fv.nome AS favorecido, t.descricao,
                  t.numero_parcela, t.total_parcelas, t.valor_centavos
           FROM fatura f JOIN cartao k ON k.id = f.cartao_id JOIN transacao t ON t.fatura_id = f.id
           JOIN plastico pl ON pl.id = t.plastico_id LEFT JOIN usuario pu ON pu.id = pl.portador_id
           LEFT JOIN categoria cat ON cat.id = t.categoria_id LEFT JOIN categoria g ON g.id = cat.pai_id
           LEFT JOIN favorecido fv ON fv.id = t.favorecido_id
           WHERE f.data_vencimento BETWEEN %s AND %s
           ORDER BY k.nome, f.data_vencimento, COALESCE(t.data_compra, t.data_competencia), t.criado_em""", (de, ate)).fetchall()

    wb = Workbook()
    ws = wb.active
    ws.title = "Lançamentos"
    D, V = "DD/MM/YYYY", "#,##0.00;[Red]-#,##0.00"
    _escrever(ws, [("Competência", 13, D), ("Caixa", 13, D), ("Data da compra", 14, D), ("Tipo", 15, None), ("Estado", 12, None), ("Valor (R$)", 14, V),
                   ("Conta", 22, None), ("Cartão", 26, None), ("Portador", 18, None), ("Grupo", 22, None), ("Categoria", 26, None), ("Favorecido", 26, None),
                   ("Descrição", 34, None), ("Forma de pagamento", 18, None), ("Parcela", 9, None), ("Lançado por", 18, None), ("Origem", 12, None)],
              ([r["data_competencia"], r["data_caixa"], r["data_compra"], r["tipo"], r["estado"], _reais(r["valor_centavos"]), r["conta"],
                _cartao(r["cartao"], r["final"], r["tipo_cartao"]), r["portador"], r["grupo"], r["categoria"], r["favorecido"], r["descricao"],
                r["forma_pagamento"], f"{r['numero_parcela']}/{r['total_parcelas']}" if r["total_parcelas"] else None, r["lancado_por"], r["origem"]]
               for r in lancs))
    ws2 = wb.create_sheet("Faturas")
    _escrever(ws2, [("Conta de cartão", 26, None), ("Mês de referência", 16, "MM/YYYY"), ("Fechamento", 13, D), ("Vencimento", 13, D),
                    ("Situação", 11, None), ("Total da fatura (R$)", 20, V), ("Itens", 8, None)],
              ([r["cartao"], r["mes_referencia"], r["data_fechamento"], r["data_vencimento"], r["status"], _reais(-r["total"]), r["itens"]] for r in faturas))
    ws3 = wb.create_sheet("Itens das faturas")
    _escrever(ws3, [("Conta de cartão", 24, None), ("Vencimento", 13, D), ("Situação", 11, None), ("Data da compra", 14, D), ("Cartão", 22, None),
                    ("Portador", 18, None), ("Grupo", 22, None), ("Categoria", 26, None), ("Favorecido", 26, None), ("Descrição", 34, None),
                    ("Parcela", 9, None), ("Valor (R$)", 14, V)],
              ([r["cartao"], r["data_vencimento"], r["status"], r["data_compra"] or r["data_competencia"], f"·· {r['final']}" + (" (virtual)" if r["tipo_cartao"] == "virtual" else ""),
                r["portador"], r["grupo"], r["categoria"], r["favorecido"], r["descricao"],
                f"{r['numero_parcela']}/{r['total_parcelas']}" if r["total_parcelas"] else None, _reais(-r["valor_centavos"])] for r in itens))
    buf = io.BytesIO()
    wb.save(buf)
    nome = f"financas_{de.isoformat()}_a_{ate.isoformat()}.xlsx"
    return Response(buf.getvalue(), media_type=XLSX, headers={"Content-Disposition": f'attachment; filename="{nome}"'})
